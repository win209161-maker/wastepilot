"""
WastePilot — Export clean customer list to Excel
Run: python export_to_excel.py
"""
import requests
import getpass
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

SUPABASE_URL = "https://fiftiotizthtwayudpkj.supabase.co"
SUPABASE_ANON_KEY = "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6ImZpZnRpb3RpenRodHdheXVkcGtqIiwicm9sZSI6ImFub24iLCJpYXQiOjE3OTA2NzA4NzgsImV4cCI6MjEwNjI0Njg3OH0.QQ1eqNdvHVOhjle-v1Gbn4Epu5WCOo7RQF0938pwN6c"
EXCEL_PATH = r"C:\Users\user\Barry_project\trash project\abonnement2026.xlsx"


def sign_in(email, password):
    r = requests.post(
        f"{SUPABASE_URL}/auth/v1/token?grant_type=password",
        json={"email": email, "password": password},
        headers={"apikey": SUPABASE_ANON_KEY, "Content-Type": "application/json"},
    )
    if r.status_code != 200:
        print(f"Login failed: {r.text}")
        return None
    return r.json()["access_token"]


def fetch_all(token, table, select="*", order=None):
    headers = {
        "apikey": SUPABASE_ANON_KEY,
        "Authorization": f"Bearer {token}",
        "Range": "0-9999",
    }
    params = {"select": select}
    if order:
        params["order"] = order
    r = requests.get(
        f"{SUPABASE_URL}/rest/v1/{table}",
        headers=headers,
        params=params,
    )
    r.raise_for_status()
    return r.json()


def update_excel(customers, sectors_map):
    wb = openpyxl.load_workbook(EXCEL_PATH)

    # Remove old "Base initiale" and recreate clean
    if "Base initiale" in wb.sheetnames:
        del wb["Base initiale"]

    ws = wb.create_sheet("Base initiale", 0)

    # Title row
    ws.merge_cells("A1:K1")
    title_cell = ws["A1"]
    title_cell.value = "LISTE DES ABONNÉS — WASTEPILOT CONAKRY"
    title_cell.font = Font(bold=True, size=13)
    title_cell.alignment = Alignment(horizontal="center")

    # Header row
    headers = ["N°", "Nom", "Prénoms", "Contacts", "Quartier", "Date demande",
               "Secteur", "Concession", "Foyer/Réf", "Numéro abonné", "Statut"]
    header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=10)

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    # Sort customers: by sector then last_name
    def sort_key(c):
        sector_name = sectors_map.get(c.get("sector_id", ""), "ZZZ")
        return (sector_name, (c.get("last_name") or "").upper(), (c.get("first_name") or "").upper())

    customers_sorted = sorted(customers, key=sort_key)

    alt_fill = PatternFill(start_color="EEF2FF", end_color="EEF2FF", fill_type="solid")

    for i, c in enumerate(customers_sorted, 1):
        row = i + 2
        sector_name = sectors_map.get(c.get("sector_id", ""), "")
        date_val = c.get("request_date") or ""

        row_data = [
            i,
            (c.get("last_name") or "").upper(),
            c.get("first_name") or "",
            c.get("phone") or "",
            c.get("neighborhood") or "",
            date_val,
            sector_name,
            c.get("concession") or "",
            c.get("reference") or c.get("address") or "",
            c.get("subscriber_id") or "",
            c.get("status") or "active",
        ]

        for col, val in enumerate(row_data, 1):
            cell = ws.cell(row=row, column=col, value=val)
            if i % 2 == 0:
                cell.fill = alt_fill

    # Column widths
    widths = [5, 20, 20, 14, 12, 14, 14, 20, 20, 18, 10]
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w

    # Freeze header
    ws.freeze_panes = "A3"

    wb.save(EXCEL_PATH)
    print(f"\n✓ Wrote {len(customers)} customers to 'Base initiale' sheet in {EXCEL_PATH}")


if __name__ == "__main__":
    print("WastePilot — Export customers to Excel")
    print("=" * 40)

    email = input("Email [admin@wastepilot.app]: ").strip() or "admin@wastepilot.app"
    password = getpass.getpass("Password: ")

    print("\nSigning in...")
    token = sign_in(email, password)
    if not token:
        exit(1)
    print("✓ Signed in")

    print("Fetching customers...")
    customers = fetch_all(token, "customers", select="*", order="last_name.asc")
    print(f"✓ {len(customers)} customers")

    print("Fetching sectors...")
    sectors = fetch_all(token, "sectors", select="id,name,code")
    sectors_map = {s["id"]: s["name"] for s in sectors}
    print(f"✓ {len(sectors)} sectors")

    update_excel(customers, sectors_map)
